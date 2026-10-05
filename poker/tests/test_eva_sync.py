from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from poker.forms import SprintCapacityForm, SprintForm
from poker.models import Project, ProjectMember, Sprint, SprintResource, SprintTask, Task, VotingSession
from poker.sprint_import import parse_sprint_file, save_sprint_import
from poker.tests.test_sprint_import import PLANNING_HEADERS, row, upload


class EvaSyncTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user('eva-sync')
        self.project = Project.objects.create(owner=self.user, name='ABS')
        self.client.force_login(self.user)

    def sync(self, rows, *, quota=True, sprints=True, xlsx=False):
        headers = (*PLANNING_HEADERS, *(('Спринты',) if sprints else ()), *(('Тип квоты',) if quota else ()))
        return save_sprint_import(self.project, parse_sprint_file(upload(rows, headers=headers, xlsx=xlsx)))

    def test_mixed_file_create_update_transfer_clear_and_idempotence(self):
        rows = [row('ABS-1', estimate=0, sprints='') + ['Новая функциональность'],
                row('ABS-2', estimate='7,5', sprints='2026.20.abs', competency='Разработка BE') + ['Технологическая']]
        first = self.sync(rows, xlsx=True)
        self.assertEqual(first.counts['created'], 2)
        self.assertEqual(first.counts['sprints_created'], 1)
        task = self.project.tasks.get(number='ABS-2')
        self.assertEqual(task.estimate, Decimal('7.50'))
        self.assertEqual(task.competency, 'development_be')
        self.assertEqual(task.quota, 'Технологическая')
        self.assertEqual(len(first.tasks), 1)
        self.assertEqual(self.sync(rows).counts['unchanged'], 2)
        old = self.project.sprints.get()
        change = row('ABS-2', estimate=20, sprints='2026.21.abs') + ['Новая функциональность']
        self.sync([change])
        task.refresh_from_db()
        self.assertEqual(task.estimate, 20)
        self.assertEqual(task.sprint_items.filter(status='planned').count(), 1)
        moved = task.sprint_items.get(sprint=old)
        self.assertEqual(moved.status, 'transferred')
        self.assertEqual(moved.transferred_to.name, '2026.21.abs')
        self.sync([row('ABS-2', estimate=20, sprints='') + ['']])
        self.assertEqual(task.sprint_items.filter(status='planned').count(), 0)
        self.assertEqual(task.sprint_items.filter(status='removed').count(), 1)
        # Reassignment reuses the old row and clears transfer metadata.
        self.sync([change])
        self.assertEqual(task.sprint_items.filter(status='planned').count(), 1)
        self.assertIsNone(task.sprint_items.get(status='planned').transferred_at)

    def test_missing_columns_preserve_but_explicit_empty_clears(self):
        self.sync([row(sprints='2026.20.abs') + ['Операционные задачи']])
        self.sync([row(estimate=32)], quota=False, sprints=False)
        task = self.project.tasks.get()
        self.assertEqual(task.quota, 'Операционные задачи')
        self.assertEqual(task.eva_sprints, '2026.20.abs')
        self.assertEqual(task.sprint_items.filter(status='planned').count(), 1)
        self.assertEqual(task.estimate, 32)
        self.sync([row(sprints='') + ['']])
        task.refresh_from_db()
        self.assertEqual((task.quota, task.eva_sprints), ('', ''))
        self.assertFalse(task.sprint_items.filter(status='planned').exists())

    def test_multi_sprint_duplicate_names_and_archived_targets_are_conflicts(self):
        self.sync([row(sprints='Original') + ['Квота']])
        task = self.project.tasks.get()
        for name in ('A; B', '["A", "B"]', '{"id": "uuid"}'):
            saved = self.sync([row(estimate=52, sprints=name) + ['Другая']])
            self.assertEqual(saved.counts['conflicts'], 1)
            task.refresh_from_db()
            self.assertEqual((task.estimate, task.quota), (12, 'Квота'))
        for _ in range(2):
            Sprint.objects.create(project=self.project, name='Duplicate')
        self.assertEqual(self.sync([row(sprints='Duplicate') + ['']]).counts['conflicts'], 1)
        target = Sprint.objects.create(project=self.project, name='Finished', status='completed')
        self.assertEqual(self.sync([row(sprints=target.name) + ['']]).counts['conflicts'], 1)
        self.assertEqual(task.sprint_items.get(status='planned').sprint.name, 'Original')

    def test_completed_plan_is_not_silently_rewritten(self):
        self.sync([row(sprints='Original') + ['Квота']])
        self.project.sprints.update(status='completed')
        result = self.sync([row(estimate=52, sprints='Next') + ['Другая']])
        self.assertEqual(result.counts['conflicts'], 1)
        task = self.project.tasks.get()
        self.assertEqual(task.estimate, 12)
        self.assertEqual(task.quota, 'Квота')
        self.assertFalse(self.project.sprints.filter(name='Next').exists())

    def test_identity_conflict_does_not_create_sprint_or_overwrite_quota(self):
        self.sync([row(sprints='') + ['Квота']])
        bad = row(sprints='New') + ['Другая']
        bad[0] = row('ABS-2')[0]
        saved = self.sync([bad])
        self.assertEqual(saved.counts['conflicts'], 1)
        self.assertFalse(self.project.sprints.exists())
        self.assertEqual(self.project.tasks.get().quota, 'Квота')

    def test_quotas_count_tasks_not_hours_and_exclude_transfer_history(self):
        self.sync([row('ABS-1', estimate=52, sprints='Plan') + ['Новая функциональность'],
                   row('ABS-2', estimate=1, sprints='Plan') + ['Новая функциональность'],
                   row('ABS-3', estimate=2, sprints='Plan') + ['Дефекты'],
                   row('ABS-4', estimate=0, sprints='Plan') + ['']])
        self.sync([row('ABS-5', sprints='Other') + ['Прочее']])
        sprint = self.project.sprints.get(name='Plan')
        quotas = {r['label']: r for r in sprint.quota_rows}
        self.assertEqual(quotas['Новая функциональность']['count'], 2)
        self.assertEqual(quotas['Новая функциональность']['percent'], 50)
        self.assertEqual(quotas['Дефекты']['percent'], 25)
        self.assertEqual(quotas['Не указана']['percent'], 25)
        self.assertEqual(sum(r['percent'] for r in quotas.values()), 100)
        self.sync([row('ABS-3', sprints='Other') + ['Дефекты']])
        self.assertEqual(sum(r['count'] for r in Sprint.objects.get(pk=sprint.pk).quota_rows), 3)
        page = self.client.get(sprint.get_absolute_url())
        self.assertContains(page, 'Распределение задач по квотам')
        self.assertContains(page, 'Новая функциональность')
        self.assertContains(page, '66,67%')

    def test_export_quota_and_formula_like_values_are_literal(self):
        self.sync([row(sprints='Plan') + ['=1+1']])
        sprint = self.project.sprints.get()
        response = self.client.get(reverse('poker:sprint_export', args=[sprint.pk]))
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(book['Квоты']['A2'].value, '=1+1')
        self.assertEqual(book['Квоты']['A2'].data_type, 's')
        self.assertEqual(book['Квоты']['B2'].value, 1)
        self.assertEqual(book['Квоты']['C2'].value, 1)
        self.assertEqual(book['Задачи спринта']['H2'].data_type, 's')
        book.close()
        response = self.client.get(reverse('poker:sprint_export_eva', args=[sprint.pk]))
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(book.active['L1'].value, 'Тип квоты')
        self.assertEqual(book.active['L2'].data_type, 's')
        book.close()

    def test_types_capacity_forms_and_copy_include_new_fields(self):
        sprint = Sprint.objects.create(project=self.project, name='Plan', development_abs_capacity=10,
                                       development_be_capacity=15, development_fe_capacity=7, defect_capacity=4)
        for i, (name, estimate) in enumerate([('Разработка АБС',12), ('Разработка BE',8), ('Разработка FE',4), ('Дефект',5)], 1):
            self.sync([row(f'ABS-{i}', sprints='Plan', competency=name, estimate=estimate) + ['']])
        sprint = Sprint.objects.get(pk=sprint.pk)
        rows = {r['key']:r for r in sprint.competency_capacity_rows}
        self.assertEqual(rows['development_abs']['overage'], 2)
        self.assertEqual(rows['development_be']['remaining'], 7)
        self.assertEqual(rows['development_fe']['remaining'], 3)
        self.assertEqual(rows['defect']['overage'], 1)
        self.assertEqual(sprint.total_estimate, 29)
        self.assertEqual(sprint.capacity_total, 36)
        for cls in (SprintForm, SprintCapacityForm):
            self.assertIn('defect_capacity', cls().fields)
            self.assertIn('development_be_capacity', cls().fields)
        self.client.post(reverse('poker:sprint_copy', args=[sprint.pk]))
        copied = self.project.sprints.exclude(pk=sprint.pk).get()
        self.assertEqual(copied.development_be_capacity, 15)
        self.assertEqual(copied.defect_capacity, 4)
        page = self.client.get(sprint.get_absolute_url())
        self.assertContains(page, 'Разработка BE')
        self.assertContains(page, 'Дефект')
        page = self.client.get(self.project.get_absolute_url() + '?competency=defect')
        self.assertEqual(len(page.context['tasks']), 1)
        self.assertContains(page, 'ABS-4')

    def test_team_capacity_is_separate_for_each_specialization(self):
        sprint = Sprint.objects.create(project=self.project, name='Team', capacity_source='team',
                                       start_date=date(2026,10,5), end_date=date(2026,10,9))
        for key in ('development_abs','development_be','development_fe','defect'):
            member = ProjectMember.objects.create(project=self.project, full_name=key, competency=key)
            SprintResource.objects.create(sprint=sprint, member=member, full_name=key, competency=key)
        self.assertEqual(sprint.team_capacities['development_abs'], 40)
        self.assertEqual(sprint.team_capacities['defect'], 40)
        self.assertEqual(sprint.team_capacities['development'], 0)
        self.assertEqual(sprint.capacity_total, 160)

    def test_project_import_reports_issues_and_does_not_touch_missing_tasks(self):
        other = Task.objects.create(project=self.project, number='LOCAL', title='Keep')
        bad = row(estimate='=1+1', sprints='Unexpected') + ['Квота']
        data = upload([bad], headers=(*PLANNING_HEADERS, 'Спринты','Тип квоты'))
        response = self.client.post(reverse('poker:task_import_file', args=[self.project.pk]), {'task_file':data}, follow=True)
        self.assertContains(response, 'некорректных оценок: 1')
        self.assertContains(response, 'строка пропущена')
        self.assertEqual(list(self.project.tasks.all()), [other])
        self.assertFalse(self.project.sprints.exists())

    def test_room_sync_updates_estimates_but_only_queues_unestimated_tasks(self):
        room = VotingSession.objects.create(project=self.project, name='Room')
        data = upload([row('ABS-1', estimate=0), row('ABS-2', estimate='7.5')])
        response = self.client.post(reverse('poker:session_import_file', args=[room.pk]), {'task_file': data}, follow=True)
        self.assertContains(response, 'В очередь добавлено: 1')
        self.assertEqual(room.queue_items.get().task.number, 'ABS-1')
        self.assertEqual(self.project.tasks.get(number='ABS-2').estimate, Decimal('7.5'))
        self.assertEqual(self.project.tasks.count(), 2)
