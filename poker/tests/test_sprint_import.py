from io import BytesIO
from uuid import UUID

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from poker.models import Participant, Project, Sprint, SprintTask, Task, Vote, VotingRound, VotingSession
from poker.sprint_import import available_sprint_tasks, parse_sprint_file, save_sprint_import
from poker.tests.test_task_import import HEADERS, csv_upload, task_row, xlsx_upload


PLANNING_HEADERS = (*HEADERS, "Статус.Имя статуса", "Кеш: Тип статуса", "Статус")


def row(number="ABS-SA-1", *, estimate=12, status="Открыта", sprints=None, **kwargs):
    result = task_row(number, estimate=estimate, **kwargs)
    # Distinct, deterministic EVA identifiers for different codes.
    result[0] = "CmfTask:" + str(UUID(int=int(number.rsplit("-", 1)[1])))
    result.extend((status, "Открыт", "Ожидание разработки"))
    if sprints is not None:
        result.append(sprints)
    return result


def upload(rows, *, xlsx=False, sprints=False, headers=None):
    headers = headers or (*PLANNING_HEADERS, *(('Спринты',) if sprints else ()))
    return (xlsx_upload if xlsx else csv_upload)(rows, headers=headers)


class SprintFileParserTests(SimpleTestCase):
    def test_csv_and_xlsx_select_competencies_positive_estimates_and_open_statuses(self):
        rows = [
            row("ABS-SA-1", estimate="12,0"),
            row("ABS-ABS-2", competency="Разработка АБС", status="В работе", estimate=32),
            row("ABS-QA-3", competency="Тестирование", status="В бэклоге", estimate=4),
            row("ABS-SA-4", estimate=0), row("ABS-SA-5", estimate=""),
            row("ABS-SA-6", status="Выполнена"), row("ABS-SA-7", status="CLOSED"),
            row("ABS-SA-8", competency="Дефект", estimate=5),
        ]
        for xlsx in (False, True):
            with self.subTest(xlsx=xlsx):
                parsed = parse_sprint_file(upload(rows, xlsx=xlsx))
                self.assertEqual([r.estimate for r in parsed.rows if not r.skip_reason], [12, 32, 4])
                self.assertEqual(parsed.counts['unestimated'], 2)
                self.assertEqual(parsed.counts['closed'], 2)
                self.assertEqual(parsed.counts['other_type'], 1)
                self.assertFalse(parsed.has_sprints_column)

    def test_does_not_confuse_unrelated_status_field_or_estimate_columns(self):
        data = row()
        data[-1] = "CLOSED"
        data[6:8] = [52, 32]
        self.assertEqual(parse_sprint_file(upload([data])).rows[0].estimate, 12)
        data[-2] = "CLOSED"
        self.assertEqual(parse_sprint_file(upload([data])).counts['closed'], 1)
        with self.assertRaisesMessage(ValidationError, 'Статус.Имя статуса'):
            parse_sprint_file(upload([data[:9]], headers=HEADERS))

    def test_invalid_estimates_are_reported_without_rounding_or_import(self):
        values = [5, -1, "NaN", "Infinity", "=12", "sNaN", False, "нет"]
        parsed = parse_sprint_file(upload([
            row(f"ABS-SA-{n}", estimate=value) for n, value in enumerate(values, 1)
        ], xlsx=True))
        self.assertEqual(parsed.counts['invalid_estimate'], len(values))
        self.assertTrue(all(item.skip_reason and item.estimate is None for item in parsed.rows))
        self.assertEqual(len(parsed.issues), len(values))

    def test_optional_eva_sprints_are_excluded_without_guessing_assignments(self):
        parsed = parse_sprint_file(upload([
            row("ABS-SA-1", sprints="Спринт 1"), row("ABS-SA-2", sprints="Нет"),
        ], sprints=True))
        self.assertTrue(parsed.has_sprints_column)
        self.assertEqual(parsed.counts['eva_assigned'], 1)
        self.assertEqual(parsed.rows[1].eva_sprints, "")

    def test_duplicate_rows_and_identifier_conflicts(self):
        parsed = parse_sprint_file(upload([row(), row()]))
        self.assertEqual(len(parsed.rows), 1)
        self.assertEqual(parsed.counts['duplicates'], 1)
        with self.assertRaisesMessage(ValidationError, 'разными данными'):
            parse_sprint_file(upload([row(), row(estimate=20)]))
        duplicate = row("ABS-SA-2")
        duplicate[0] = row()[0]
        with self.assertRaisesMessage(ValidationError, 'разных кодов'):
            parse_sprint_file(upload([row(), duplicate]))


class SprintPlanningTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_user("planner", password="secret")

    def setUp(self):
        self.project = Project.objects.create(owner=self.user, name="ABS")
        self.sprint = Sprint.objects.create(project=self.project, name="Спринт 1", analysis_capacity=20,
                                           development_capacity=40, testing_capacity=8)
        self.next_sprint = Sprint.objects.create(project=self.project, name="Спринт 2")
        self.client.force_login(self.user)

    def load(self, rows, **kwargs):
        return save_sprint_import(self.project, parse_sprint_file(upload(rows, **kwargs)))

    def test_full_import_plan_remaining_tasks_next_sprint_and_export_workflow(self):
        response = self.client.post(reverse('poker:sprint_import', args=[self.sprint.pk]), {
            'task_file': upload([row(), row("ABS-ABS-2", competency="Разработка АБС", estimate=32),
                                 row("ABS-QA-3", competency="Тестирование", estimate=4)], xlsx=True),
        })
        self.assertRedirects(response, self.sprint.get_absolute_url() + '#available-tasks')
        self.assertEqual(self.project.tasks.count(), 3)
        self.assertEqual(Vote.objects.count(), 0)
        self.assertEqual(VotingRound.objects.count(), 0)
        first = self.project.tasks.get(number="ABS-SA-1")
        self.assertIsNone(first.estimate_count)
        self.assertEqual(first.estimate, 12)
        self.assertIsNone(first.average_estimate)
        self.client.post(reverse('poker:sprint_add_tasks', args=[self.sprint.pk]), {'task_ids': [first.pk]})
        page = self.client.get(self.next_sprint.get_absolute_url())
        self.assertEqual(len(page.context['available_tasks']), 2)
        self.assertNotIn(first, page.context['available_tasks'])
        # Even a stale or crafted request cannot plan the task twice.
        self.client.post(reverse('poker:sprint_add_tasks', args=[self.next_sprint.pk]), {'task_ids': [first.pk]})
        self.assertEqual(first.sprint_items.filter(status='planned').count(), 1)
        saved = self.load([row(estimate=32), row("ABS-ABS-2", competency="Разработка АБС", estimate=32)])
        self.assertEqual(saved.counts['planned_or_completed'], 1)
        self.assertEqual(saved.counts['unchanged'], 1)
        first.refresh_from_db()
        self.assertEqual(first.estimate, 12)
        self.assertEqual(Sprint.objects.get(pk=self.sprint.pk).total_estimate, 12)
        export = self.client.get(reverse('poker:sprint_export_eva', args=[self.sprint.pk]))
        book = load_workbook(BytesIO(export.content))
        sheet = book['План для EVA']
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet['A2'].value, first.number)
        self.assertEqual(sheet['D2'].value, 12)
        self.assertEqual(sheet['E2'].value, first.eva_identifier)
        self.assertEqual(sheet['G2'].value, self.sprint.name)
        self.assertEqual(sheet['K2'].value, self.sprint.pk)
        book.close()
        self.client.post(reverse('poker:sprint_remove_task', args=[self.sprint.pk, first.pk]))
        self.assertIn(first, available_sprint_tasks(self.project))

    def test_repeat_import_updates_only_free_eva_estimates(self):
        self.load([row()])
        original = self.project.tasks.get()
        self.load([row(estimate=20, title='Уточнённая задача')])
        original.refresh_from_db()
        self.assertEqual(self.project.tasks.count(), 1)
        self.assertEqual(original.estimate, 20)
        self.assertEqual(original.title, 'Уточнённая задача')
        self.assertIsNone(original.estimate_sum)

    def test_existing_votes_estimates_and_conflicts_are_preserved(self):
        task = Task.objects.create(project=self.project, number='ABS-SA-1', title='Согласованная',
                                   status='estimated', estimate_sum=24, estimate_count=2)
        saved = self.load([row(estimate=20)])
        self.assertEqual(saved.counts['conflicts'], 1)
        task.refresh_from_db()
        self.assertEqual((task.estimate_sum, task.estimate_count, task.title), (24, 2, 'Согласованная'))
        self.load([row()])
        task.refresh_from_db()
        self.assertIsNone(task.imported_estimate)
        self.assertEqual(task.estimate_count, 2)

    def test_real_zero_vote_result_stays_available_when_eva_exports_zero(self):
        task = Task.objects.create(project=self.project, number='ABS-SA-1', title='Ноль',
                                   status='estimated', estimate_sum=0, estimate_count=2)
        self.load([row(estimate=0)])
        self.assertIn(task, available_sprint_tasks(self.project))

    def test_closed_assigned_and_unestimated_snapshots_remove_previously_available_tasks(self):
        for status, estimate, sprints in [('Закрыта', 12, ''), ('Открыта', 12, 'Другой'), ('Открыта', 0, '')]:
            with self.subTest(status=status, estimate=estimate, sprints=sprints):
                self.load([row(sprints='')], sprints=True)
                self.load([row(status=status, estimate=estimate, sprints=sprints)], sprints=True)
                self.assertEqual(available_sprint_tasks(self.project).count(), 0)

    def test_missing_sprint_column_does_not_clear_known_eva_assignment(self):
        self.load([row()])
        self.load([row(sprints='Уже назначенный спринт')], sprints=True)
        saved = self.load([row()])
        self.assertEqual(saved.counts['eva_assigned'], 1)
        self.assertEqual(available_sprint_tasks(self.project).count(), 0)
        self.load([row(sprints='')], sprints=True)
        self.assertEqual(available_sprint_tasks(self.project).count(), 1)

    def test_fresh_snapshot_can_unblock_plan_without_changing_reserved_hours(self):
        self.load([row()])
        task = self.project.tasks.get()
        SprintTask.objects.create(sprint=self.sprint, task=task)
        self.load([row(status='Закрыта')])
        self.assertEqual(self.client.get(reverse('poker:sprint_export_eva', args=[self.sprint.pk])).status_code, 302)
        self.load([row(estimate=20)])
        task.refresh_from_db()
        self.assertEqual(task.estimate, 12)
        self.assertFalse(task.eva_unavailable)
        self.assertEqual(self.client.get(reverse('poker:sprint_export_eva', args=[self.sprint.pk])).status_code, 200)

    def test_active_voting_is_not_overwritten(self):
        task = Task.objects.create(project=self.project, number='ABS-SA-1', title='Голосуют')
        session = VotingSession.objects.create(project=self.project, name='Оценка')
        VotingRound.objects.create(session=session, task=task, number=1)
        saved = self.load([row()])
        self.assertEqual(saved.counts['voting'], 1)
        task.refresh_from_db()
        self.assertIsNone(task.estimate)

    def test_new_vote_acceptance_replaces_imported_estimate_without_fake_votes(self):
        self.load([row()])
        self.load([row(estimate=0)])
        task = self.project.tasks.get()
        session = VotingSession.objects.create(project=self.project, name='Оценка')
        round_ = VotingRound.objects.create(session=session, task=task, number=1)
        person = Participant.objects.create(session=session, name='Участник')
        Vote.objects.create(voting_round=round_, participant=person, value=20)
        task.capture_estimate(round_)
        task.refresh_from_db()
        self.assertIsNone(task.imported_estimate)
        self.assertEqual(task.estimate, 20)
        self.assertEqual(task.estimate_count, 1)
        self.assertIn(task, available_sprint_tasks(self.project))

    def test_identifier_conflict_does_not_rebind_existing_task(self):
        self.load([row()])
        altered = row(estimate=32)
        altered[0] = 'CmfTask:' + str(UUID(int=999))
        self.assertEqual(self.load([altered]).counts['conflicts'], 1)
        self.assertEqual(self.project.tasks.get().estimate, 12)

    def test_bad_row_rolls_back_whole_import_and_shows_form_error(self):
        bad = row('ABS-SA-2')
        bad[0] = 'invalid'
        response = self.client.post(reverse('poker:sprint_import', args=[self.sprint.pk]), {'task_file': upload([row(), bad])})
        self.assertEqual(response.status_code, 400)
        self.assertContains(response, 'CmfTask:UUID', status_code=400)
        self.assertEqual(self.project.tasks.count(), 0)

    def test_owner_only_and_closed_sprint_mutation_guards(self):
        stranger = get_user_model().objects.create_user('other')
        self.client.force_login(stranger)
        for route in ('sprint_import', 'sprint_add_tasks', 'sprint_export_eva'):
            url = reverse('poker:' + route, args=[self.sprint.pk])
            response = self.client.get(url) if route.endswith('eva') else self.client.post(url)
            self.assertEqual(response.status_code, 404)
        self.client.force_login(self.user)
        for fields in ({'status': 'completed'}, {'status': 'planning', 'archived_at': timezone.now()}):
            Sprint.objects.filter(pk=self.sprint.pk).update(**fields)
            response = self.client.post(reverse('poker:sprint_import', args=[self.sprint.pk]), {'task_file': upload([row()])})
            self.assertEqual(response.status_code, 302)
            self.assertEqual(self.project.tasks.count(), 0)

    def test_completed_sprint_reserves_tasks_for_future_sprints(self):
        self.load([row()])
        task = self.project.tasks.get()
        SprintTask.objects.create(sprint=self.sprint, task=task)
        self.sprint.status = 'completed'
        self.sprint.save()
        self.assertEqual(available_sprint_tasks(self.project).count(), 0)

    def test_export_literal_text_and_transfer_excludes_old_membership(self):
        self.load([row(title='=HYPERLINK("http://example.test")')])
        task = self.project.tasks.get()
        SprintTask.objects.create(sprint=self.sprint, task=task)
        self.client.post(reverse('poker:sprint_transfer_tasks', args=[self.sprint.pk]), {
            'task_ids': [task.pk], 'target_sprint': self.next_sprint.pk,
        })
        self.assertEqual(task.sprint_items.filter(status='planned').count(), 1)
        old = self.client.get(reverse('poker:sprint_export_eva', args=[self.sprint.pk]))
        self.assertEqual(old.status_code, 302)
        exported = self.client.get(reverse('poker:sprint_export_eva', args=[self.next_sprint.pk]))
        book = load_workbook(BytesIO(exported.content), data_only=False)
        self.assertEqual(book.active['B2'].data_type, 's')
        self.assertEqual(book.active['G2'].value, self.next_sprint.name)
        book.close()

    def test_export_refuses_legacy_duplicate_assignments(self):
        self.load([row()])
        task = self.project.tasks.get()
        for sprint in (self.sprint, self.next_sprint):
            SprintTask.objects.create(sprint=sprint, task=task)
        response = self.client.get(reverse('poker:sprint_export_eva', args=[self.sprint.pk]))
        self.assertEqual(response.status_code, 302)
