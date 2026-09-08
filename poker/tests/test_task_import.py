import csv
from io import BytesIO, StringIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, SimpleTestCase, TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook

from poker.models import Participant, Project, Task, Vote, VotingRound, VotingSession, VotingSessionTask
from poker.task_import import EVA_TASK_URL_PREFIX, parse_task_file


HEADERS = (
    "Идентификатор объекта", "Код", "Наименование", "Текст", "Текст без html",
    "Логический тип.Имя логического типа", "Оценка", "Story Point", "Оценка задачи, час",
)
IDENTIFIER = "CmfTask:5d7adfd4-a87a-11f1-8708-9e913e917b69"


def task_row(number="ABS-1", *, estimate="", competency="Системный анализ", title="Проверить комиссию", description="<p>Первая строка</p><p>Вторая строка</p>"):
    return [IDENTIFIER, number, title, description, "Запасной текст", competency, "", "", estimate]


def csv_upload(rows, *, headers=HEADERS, encoding="utf-8-sig", delimiter=";", prefix=""):
    output = StringIO(newline="")
    output.write(prefix)
    writer = csv.writer(output, delimiter=delimiter, lineterminator="\r\n")
    writer.writerow(headers)
    writer.writerows(rows)
    return SimpleUploadedFile("eva.csv", output.getvalue().encode(encoding), content_type="text/csv")


def xlsx_upload(rows, *, headers=HEADERS):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Выгрузка"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return SimpleUploadedFile("eva.xlsx", output.getvalue())


class TaskFileParserTests(SimpleTestCase):
    def test_csv_encodings_delimiters_and_multiline_description(self):
        for encoding, delimiter, prefix in (
            ("utf-8-sig", ";", ""), ("cp1251", ",", ""), ("utf-16", "\t", "sep=\t\r\n"),
        ):
            with self.subTest(encoding=encoding, delimiter=delimiter):
                row = task_row(title='Отчёт; "комиссия", часть 1', description="<p>Строка 1</p>\n<p>Строка 2</p>")
                parsed = parse_task_file(csv_upload([row], encoding=encoding, delimiter=delimiter, prefix=prefix))
                self.assertEqual(parsed.total_rows, 1)
                self.assertEqual(parsed.tasks[0].title, row[2])
                self.assertEqual(parsed.tasks[0].description, "Строка 1\nСтрока 2")
                self.assertEqual(parsed.tasks[0].external_url, EVA_TASK_URL_PREFIX + IDENTIFIER)

    def test_xlsx_skips_every_nonempty_hour_estimate_including_zero_and_formula(self):
        estimates = [0, "0", "0.0", "0,00", 32, "5", "—", "=1-1", "неизвестно", None, "", "  "]
        rows = [task_row(f"ABS-{i}", estimate=value) for i, value in enumerate(estimates)]
        parsed = parse_task_file(xlsx_upload(rows))
        self.assertEqual(parsed.skipped_estimated, 9)
        self.assertEqual(parsed.skipped_zero, 4)
        self.assertEqual([task.number for task in parsed.tasks], ["ABS-9", "ABS-10", "ABS-11"])

    def test_uses_hour_field_and_exact_headers_in_any_order(self):
        row = task_row()
        row[6:8] = [52, 32]  # Unrelated EVA estimates do not decide eligibility.
        indexes = [8, 2, 0, 5, 1, 3, 4, 7, 6]
        headers = ["  " + HEADERS[i].upper() + "  " for i in indexes]
        parsed = parse_task_file(csv_upload([[row[i] for i in indexes]], headers=headers))
        self.assertEqual(len(parsed.tasks), 1)
        missing = list(HEADERS)
        missing[-1] = "Предварительная оценка"
        with self.assertRaisesMessage(ValidationError, "Оценка задачи, час"):
            parse_task_file(csv_upload([row], headers=missing))

    def test_maps_eva_competencies_without_inferring_from_title(self):
        expected = {
            "Системный анализ": "analysis", "Аналитика": "analysis",
            "Разработка АБС": "development", "Разработка BE": "development",
            "Разработка FE": "development", "Разработка Битрикс": "development",
            "Тестирование": "testing", "Тестирование - на DEV": "testing",
        }
        rows = [task_row(f"ABS-{i}", competency=name, title="[SA] Общая задача") for i, name in enumerate(expected)]
        self.assertEqual([task.competency for task in parse_task_file(csv_upload(rows)).tasks], list(expected.values()))

    def test_description_converts_html_preserves_paragraphs_and_ignores_scripts(self):
        row = task_row(description='<p>Первый &amp; второй</p><script>alert(1)</script><style>body{}</style><ul><li>Три</li><li>Четыре<br>Пять</li></ul>')
        self.assertEqual(parse_task_file(csv_upload([row])).tasks[0].description, "Первый & второй\nТри\nЧетыре\nПять")
        row[3] = ""
        row[4] = "Текст без HTML\nСледующая строка"
        self.assertEqual(parse_task_file(csv_upload([row])).tasks[0].description, row[4])

    def test_exact_duplicate_is_counted_and_conflicting_duplicate_rejected(self):
        row = task_row()
        parsed = parse_task_file(csv_upload([row, row]))
        self.assertEqual(len(parsed.tasks), 1)
        self.assertEqual(parsed.duplicates, 1)
        with self.assertRaisesMessage(ValidationError, "повторяется с разными данными"):
            parse_task_file(csv_upload([row, task_row(title="Другое название")]))

    def test_estimated_duplicate_cannot_be_reimported_as_blank(self):
        for rows in ([task_row(), task_row(estimate=0)], [task_row(estimate=32), task_row()]):
            with self.subTest(rows=rows):
                parsed = parse_task_file(csv_upload(rows))
                self.assertEqual(parsed.tasks, [])
                self.assertEqual(parsed.skipped_estimated, 2)

    def test_unknown_type_and_invalid_identifier_give_row_errors(self):
        for row, error in ((task_row(competency="Дефект"), "неизвестный тип"), (task_row(), "идентификатор объекта")):
            if error == "идентификатор объекта":
                row[0] = 'javascript:alert(1)'
            with self.subTest(error=error), self.assertRaisesMessage(ValidationError, "Строка 2: " + ("нужен " if error == "идентификатор объекта" else "") + error):
                parse_task_file(csv_upload([row]))
        # Irrelevant data on rows that must be skipped does not block the file.
        self.assertEqual(parse_task_file(csv_upload([task_row(competency="Дефект", estimate=5)])).skipped_estimated, 1)

    def test_bad_files_and_limits_fail_with_validation_errors(self):
        for name, content in (("file.xls", b"bad"), ("file.xlsx", b"bad"), ("file.csv", b"")):
            with self.subTest(name=name), self.assertRaises(ValidationError):
                parse_task_file(SimpleUploadedFile(name, content))
        with patch("poker.task_import.MAX_FILE_BYTES", 4), self.assertRaisesMessage(ValidationError, "10 МБ"):
            parse_task_file(csv_upload([task_row()]))
        with patch("poker.task_import.MAX_ROWS", 1), self.assertRaisesMessage(ValidationError, "1 строк"):
            parse_task_file(csv_upload([task_row(), task_row("ABS-2")]))

    def test_xlsx_selects_data_sheet_and_rejects_ambiguous_sheets(self):
        workbook = Workbook()
        workbook.active.append(["Справка"])
        sheet = workbook.create_sheet("Задачи")
        sheet.append(HEADERS)
        sheet.append(task_row())
        output = BytesIO()
        workbook.save(output)
        self.assertEqual(len(parse_task_file(SimpleUploadedFile("eva.xlsx", output.getvalue())).tasks), 1)
        workbook.copy_worksheet(sheet)
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        with self.assertRaisesMessage(ValidationError, "несколько листов"):
            parse_task_file(SimpleUploadedFile("eva.xlsx", output.getvalue()))


class TaskFileImportViewsTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.owner = get_user_model().objects.create_user("import-owner")
        cls.project = Project.objects.create(owner=cls.owner, name="Импорт EVA")

    def setUp(self):
        self.client.force_login(self.owner)
        self.url = reverse("poker:task_import_file", args=[self.project.pk])

    def test_project_import_skips_source_estimates_and_is_idempotent(self):
        rows = [task_row(), task_row("ABS-2", estimate=0), task_row("ABS-3", estimate=32)]
        response = self.client.post(self.url, {"task_file": xlsx_upload(rows)}, follow=True)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Добавлено: 1")
        self.assertContains(response, "Пропущено с оценкой в EVA: 2 (из них с нулевой: 1)")
        task = self.project.tasks.get()
        self.assertEqual(task.number, "ABS-1")
        self.assertEqual(task.competency, "analysis")
        self.assertEqual(task.description, "Первая строка\nВторая строка")
        self.assertIsNone(task.estimate)
        self.assertIsNone(task.estimate_sum)
        self.assertIsNone(task.estimate_count)
        self.assertEqual(task.external_url, EVA_TASK_URL_PREFIX + IDENTIFIER)
        response = self.client.post(self.url, {"task_file": csv_upload(rows)}, follow=True)
        self.assertContains(response, "без изменений: 1")
        self.assertEqual(self.project.tasks.count(), 1)

    def test_import_updates_only_unestimated_tasks_and_preserves_existing_estimates(self):
        pending = Task.objects.create(project=self.project, number="ABS-1", title="Старое")
        zero = Task.objects.create(project=self.project, number="ABS-2", title="Нулевая", estimate_sum=0, estimate_count=4, status=Task.Status.ESTIMATED)
        estimated = Task.objects.create(project=self.project, number="ABS-3", title="Оценена", estimate_sum=116, estimate_count=4, status=Task.Status.ESTIMATED)
        completed = Task.objects.create(project=self.project, number="ABS-4", title="Закрытая", completed_at=timezone.now())
        rows = [task_row(f"ABS-{i}", competency="Разработка АБС") for i in range(1, 5)]
        response = self.client.post(self.url, {"task_file": csv_upload(rows)}, follow=True)
        self.assertContains(response, "обновлено: 1")
        self.assertContains(response, "уже оценённых или завершённых в score-it: 3")
        pending.refresh_from_db()
        self.assertEqual(pending.competency, "development")
        self.assertEqual(pending.title, rows[0][2])
        for task, title in ((zero, "Нулевая"), (estimated, "Оценена"), (completed, "Закрытая")):
            task.refresh_from_db()
            self.assertEqual(task.title, title)
            self.assertEqual(task.external_url, "")
        self.assertEqual((zero.estimate_sum, zero.estimate_count), (0, 4))
        self.assertEqual((estimated.estimate_sum, estimated.estimate_count), (116, 4))

    def test_invalid_row_prevents_partial_import_and_reports_form_error(self):
        bad = task_row("ABS-2")
        bad[0] = "wrong"
        response = self.client.post(self.url, {"task_file": csv_upload([task_row(), bad])})
        self.assertEqual(response.status_code, 400)
        self.assertContains(response, "Строка 3:", status_code=400)
        self.assertFalse(self.project.tasks.exists())

    def test_import_requires_project_ownership_and_post(self):
        self.assertEqual(self.client.get(self.url).status_code, 405)
        self.client.logout()
        self.assertEqual(self.client.post(self.url, {"task_file": csv_upload([task_row()])}).status_code, 302)
        other = get_user_model().objects.create_user("other-importer")
        self.client.force_login(other)
        self.assertEqual(self.client.post(self.url, {"task_file": csv_upload([task_row()])}).status_code, 404)
        self.assertFalse(self.project.tasks.exists())

    def test_import_to_active_room_opens_new_rounds_preserves_votes_and_resets_completion(self):
        room = VotingSession.objects.create(project=self.project, name="Оценка", status=VotingSession.Status.ACTIVE, minimum_participants=1)
        task = Task.objects.create(project=self.project, number="OLD-1", title="Существующая")
        voting_round = VotingRound.objects.create(session=room, task=task)
        VotingSessionTask.objects.create(session=room, task=task, position=1, current_round=voting_round, status=VotingSessionTask.Status.ACTIVE)
        room.current_task = task
        room.save()
        participant = Participant.objects.create(session=room, name="Тест", completed_at=timezone.now())
        vote = Vote.objects.create(voting_round=voting_round, participant=participant, value=12)
        url = reverse("poker:session_import_file", args=[room.pk])
        rows = [task_row(), task_row("ABS-2", estimate=0)]
        response = self.client.post(url, {"task_file": csv_upload(rows)}, follow=True)
        self.assertContains(response, "В очередь добавлено: 1")
        self.assertEqual(room.queue_items.count(), 2)
        imported = room.queue_items.get(position=2)
        self.assertEqual(imported.task.number, "ABS-1")
        self.assertEqual(imported.current_round.status, VotingRound.Status.VOTING)
        participant.refresh_from_db()
        self.assertIsNone(participant.completed_at)
        vote.refresh_from_db()
        self.assertEqual(vote.value, 12)
        self.client.post(url, {"task_file": csv_upload(rows)})
        self.assertEqual(room.queue_items.count(), 2)
        self.assertEqual(room.rounds.count(), 2)

    def test_room_import_cannot_cross_owner_or_write_to_finished_room(self):
        room = VotingSession.objects.create(project=self.project, name="Закончена", status=VotingSession.Status.FINISHED)
        url = reverse("poker:session_import_file", args=[room.pk])
        self.assertEqual(self.client.post(url, {"task_file": csv_upload([task_row()])}).status_code, 302)
        self.assertFalse(self.project.tasks.exists())
        self.assertFalse(room.queue_items.exists())
        other = get_user_model().objects.create_user("other-room-importer")
        self.client.force_login(other)
        self.assertEqual(self.client.post(url, {"task_file": csv_upload([task_row()])}).status_code, 404)

    def test_empty_draft_room_offers_upload_and_validates_file(self):
        room = VotingSession.objects.create(project=self.project, name="Новая")
        self.assertContains(self.client.get(room.get_absolute_url()), "Загрузить и добавить в очередь")
        url = reverse("poker:session_import_file", args=[room.pk])
        self.assertEqual(self.client.post(url, {}).status_code, 400)
        self.client.post(url, {"task_file": csv_upload([task_row()])})
        queue_item = room.queue_items.get()
        self.assertIsNone(queue_item.current_round)
        self.assertEqual(queue_item.status, VotingSessionTask.Status.PENDING)

    def test_participant_receives_imported_details_before_and_after_reveal(self):
        room = VotingSession.objects.create(project=self.project, name="Карточка", status=VotingSession.Status.ACTIVE, minimum_participants=1)
        self.client.post(reverse("poker:session_import_file", args=[room.pk]), {"task_file": csv_upload([task_row()])})
        participant = Client()
        participant.post(reverse("poker:room_join", args=[room.public_token]), {"name": "Участник"})
        state_url = reverse("poker:room_state", args=[room.public_token])
        for status in (VotingRound.Status.VOTING, VotingRound.Status.REVEALED, VotingRound.Status.CLOSED):
            with self.subTest(status=status):
                room.rounds.update(status=status)
                data = participant.get(state_url).json()
                self.assertEqual(data["current_task"]["number"], "ABS-1")
                self.assertEqual(data["current_task"]["description"], "Первая строка\nВторая строка")
                self.assertEqual(data["current_task"]["external_url"], EVA_TASK_URL_PREFIX + IDENTIFIER)
                self.assertEqual(data["current_task"]["competency_label"], "Аналитика")
        page = participant.get(room.get_public_url())
        self.assertContains(page, 'id="room-task-description"')
        self.assertContains(page, 'id="result-task-description"')
        self.assertContains(page, 'target="_blank" rel="noopener noreferrer"', count=2)

    def test_organizer_escapes_plain_description(self):
        room = VotingSession.objects.create(project=self.project, name="Описание", status=VotingSession.Status.ACTIVE)
        row = task_row(description="")
        row[4] = '<img src=x onerror="alert(1)"> & текст'
        self.client.post(reverse("poker:session_import_file", args=[room.pk]), {"task_file": csv_upload([row])})
        response = self.client.get(room.get_absolute_url())
        self.assertContains(response, "&lt;img")
        self.assertNotContains(response, '<img src=x')
