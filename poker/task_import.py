"""Bounded CSV/XLSX readers and shared EVA column/type mapping."""

import csv
import re
from dataclasses import dataclass
from html.parser import HTMLParser
from io import BytesIO, StringIO
from itertools import chain
from pathlib import Path
from pathlib import PurePosixPath
import stat
from uuid import UUID
from zipfile import BadZipFile, ZipFile, ZIP_STORED, ZIP_DEFLATED

from django.core.exceptions import ValidationError
from defusedxml.ElementTree import iterparse
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .models import Task


EVA_TASK_URL_PREFIX = "https://projects.cifra.works/project/Kanban/KBB-000451?popup="
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLUMNS = 512
MAX_UNPACKED_BYTES = 64 * 1024 * 1024
MAX_ZIP_MEMBERS = 1024
MAX_SHEETS = 32
MAX_XML_NODES = 2_000_000
MAX_CELL_CHARS = 131072
HOUR_ESTIMATE_HEADER = "Оценка задачи, час"

# Deliberately exclude EVA's unrelated fields «Оценка» and «Story Point».
HEADERS = {
    "eva_parent_title": ("Родительская задача.Наименование",),
    "eva_project_name": ("Проект.Имя объекта",),
    "eva_epic_title": ("Epic.Наименование",),
    "quota": ("Тип квоты",),
    "number": ("Код", "Код задачи", "Номер задачи", "Номер"),
    "title": ("Наименование", "Название задачи", "Название"),
    "competency": ("Логический тип.Имя логического типа", "Логический тип", "Тип задачи"),
    "identifier": ("Идентификатор объекта",),
    "estimate": (HOUR_ESTIMATE_HEADER, "Оценка задачи, ч", "Оценка задачи, часы"),
    "html_description": ("Текст", "Описание"),
    "plain_description": ("Текст без html", "Описание без html"),
    # EVA also exports an unrelated bare «Статус» column. Never use it here.
    "eva_status": ("Статус.Имя статуса", "Статус задачи"),
    "eva_status_type": ("Кеш: Тип статуса", "Кэш: Тип статуса", "Тип статуса"),
    "eva_sprints": ("Спринты", "Спринт", "Спринт.Название", "Спринт.Наименование"),
}
COMPETENCIES = {
    "системный анализ": Task.Competency.ANALYSIS,
    "бизнес-анализ": Task.Competency.ANALYSIS,
    "бизнес анализ": Task.Competency.ANALYSIS,
    "аналитика": Task.Competency.ANALYSIS,
    "анализ": Task.Competency.ANALYSIS,
    "analysis": Task.Competency.ANALYSIS,
    "разработка": Task.Competency.DEVELOPMENT,
    "разработка абс": Task.Competency.DEVELOPMENT_ABS,
    "разработка битрикс": Task.Competency.DEVELOPMENT,
    "разработка be": Task.Competency.DEVELOPMENT_BE,
    "разработка fe": Task.Competency.DEVELOPMENT_FE,
    "development": Task.Competency.DEVELOPMENT,
    "dev": Task.Competency.DEVELOPMENT,
    "тестирование": Task.Competency.TESTING,
    "тестирование - на dev": Task.Competency.TESTING,
    "тестирование - на test": Task.Competency.TESTING,
    "testing": Task.Competency.TESTING,
    "qa": Task.Competency.TESTING,
    "дефект": Task.Competency.DEFECT,
    "defect": Task.Competency.DEFECT,
    "bug": Task.Competency.DEFECT,
    "": Task.Competency.NONE,
    "без типа": Task.Competency.NONE,
}


def _text(value):
    return "" if value is None else str(value).strip()


def _normalise(value):
    return " ".join(_text(value).lstrip("\ufeff").split()).casefold()


class _DescriptionText(HTMLParser):
    """Keep paragraphs and lists; never store executable source markup."""

    blocks = {
        "p", "div", "br", "li", "ul", "ol", "tr",
        "h1", "h2", "h3", "h4", "blockquote",
    }
    ignored = {"script", "style"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.ignore_depth = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.ignored:
            self.ignore_depth += 1
        if not self.ignore_depth and tag in self.blocks:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self.ignored:
            self.ignore_depth = max(0, self.ignore_depth - 1)
        elif not self.ignore_depth and tag in self.blocks:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.ignore_depth:
            self.parts.append(data)


def description_text(value):
    parser = _DescriptionText()
    parser.feed(_text(value))
    parser.close()
    lines = [line.strip() for line in "".join(parser.parts).replace("\xa0", " ").splitlines()]
    return "\n".join(line for line in lines if line)


@dataclass(frozen=True)
class ImportedTask:
    number: str
    title: str
    competency: str
    description: str
    external_url: str
    quota: str = ""
    eva_parent_title: str = ""
    eva_project_name: str = ""
    eva_epic_title: str = ""


def _column_mapping(headers):
    mapping = {}
    normalised = [_normalise(value) for value in headers]
    for key, aliases in HEADERS.items():
        for alias in aliases:
            indexes = [
                i for i, value in enumerate(normalised)
                if value == _normalise(alias)
            ]
            if len(indexes) > 1:
                raise ValidationError(f"Колонка «{alias}» встречается несколько раз.")
            if indexes:
                mapping[key] = indexes[0]
                break
    required = ("number", "title", "competency", "identifier", "estimate")
    missing = [HEADERS[key][0] for key in required if key not in mapping]
    if not {"html_description", "plain_description"} & mapping.keys():
        missing.append("Текст или Текст без html")
    if missing:
        raise ValidationError(
            "Не найдены колонки: " + ", ".join(missing)
            + ". Используйте выгрузку EVA «все поля»."
        )
    return mapping


def _csv_rows(data):
    csv.field_size_limit(MAX_CELL_CHARS)
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        text = data.decode("utf-16")
    else:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("cp1251")
    text = text.lstrip("\ufeff")
    stream = StringIO(text, newline='')
    first = stream.readline()
    delimiter = None
    if re.fullmatch(r"sep=[;,\t]", first.rstrip('\r\n'), re.IGNORECASE):
        delimiter = first.rstrip('\r\n')[-1]
    else:
        stream.seek(0)
    if delimiter is None:
        # A comma inside «Оценка задачи, час» is not a CSV delimiter.
        # Match known headings rather than relying on Sniffer's comma preference.
        header = ''
        for _ in range(MAX_ROWS + 1):
            line = stream.readline()
            if not line or line.strip():
                header = line.rstrip('\r\n')
                break
        else:
            raise ValidationError(f"Загрузите не более {MAX_ROWS} строк за один раз.")
        stream.seek(0)
        known_headers = {
            _normalise(alias) for aliases in HEADERS.values() for alias in aliases
        }
        candidates = []
        for candidate in (",", ";", "\t"):
            fields = next(csv.reader([header], delimiter=candidate))
            matches = sum(_normalise(value) in known_headers for value in fields)
            candidates.append((matches, len(fields), candidate))
        matches, field_count, delimiter = max(candidates)
        if not matches and field_count < 2:
            raise ValidationError(
                "Не удалось определить разделитель CSV: используйте запятую, "
                "точку с запятой или табуляцию."
            )
    reader = csv.reader(stream, delimiter=delimiter, strict=True)
    return _bounded_rows(reader)


def _bounded_rows(rows):
    result = []
    for row in rows:
        if len(row) > MAX_COLUMNS:
            raise ValidationError(f"В файле больше {MAX_COLUMNS} колонок.")
        result.append(tuple(row))
        if len(result) > MAX_ROWS + 1:
            raise ValidationError(f"Загрузите не более {MAX_ROWS} строк за один раз.")
    return result


def _validate_xlsx(data):
    total = 0
    nodes = 0
    class BoundedMember:
        def __init__(self, source):
            self.source = source

        def read(self, size=-1):
            nonlocal total
            chunk = self.source.read(min(size, 65536) if size >= 0 else 65536)
            total += len(chunk)
            if total > MAX_UNPACKED_BYTES:
                raise ValidationError('Содержимое XLSX слишком велико. Разделите выгрузку на несколько файлов.')
            return chunk

    with ZipFile(BytesIO(data)) as archive:
        members = archive.infolist()
        if len(members) > MAX_ZIP_MEMBERS or len({item.filename for item in members}) != len(members):
            raise ValidationError('Некорректный XLSX: слишком много или несколько одинаковых компонентов.')
        if sum(item.file_size for item in members) > MAX_UNPACKED_BYTES:
            raise ValidationError('Содержимое XLSX слишком велико. Разделите выгрузку на несколько файлов.')
        for item in members:
            path = PurePosixPath(item.filename)
            if (item.flag_bits & 1 or item.compress_type not in (ZIP_STORED, ZIP_DEFLATED)
                    or path.is_absolute() or '..' in path.parts or '\\' in item.filename
                    or ':' in item.filename or stat.S_ISLNK(item.external_attr >> 16)):
                raise ValidationError('Некорректный XLSX: неподдерживаемый или небезопасный компонент.')
            if item.is_dir():
                continue
            with archive.open(item) as source:
                stream = BoundedMember(source)
                if not item.filename.lower().endswith(('.xml', '.rels')):
                    while stream.read(65536):
                        pass
                    continue
                stack = []
                for event, element in iterparse(stream, events=('start', 'end'),
                        forbid_dtd=True, forbid_entities=True, forbid_external=True):
                    if event == 'start':
                        stack.append(element)
                        nodes += 1
                        if len(stack) > 64 or nodes > MAX_XML_NODES:
                            raise ValidationError('Структура XLSX слишком сложна. Разделите выгрузку.')
                    else:
                        if len(element.text or '') > MAX_CELL_CHARS or len(element.tail or '') > MAX_CELL_CHARS:
                            raise ValidationError('В XLSX слишком длинное значение ячейки.')
                        stack.pop()
                        element.clear()
                        if stack:
                            stack[-1].remove(element)


def _xlsx_rows(data):
    _validate_xlsx(data)
    # Keep formulas visible: an estimate formula without a cached value is filled.
    workbook = load_workbook(
        BytesIO(data), read_only=True, data_only=False, keep_links=False,
    )
    try:
        if len(workbook.worksheets) > MAX_SHEETS:
            raise ValidationError('В XLSX слишком много листов. Оставьте лист с задачами EVA.')
        candidates = []
        errors = []
        for sheet in workbook:
            # Dimension metadata is attacker-controlled and can pad small rows
            # into millions of empty cells. Bound actual rows and columns.
            sheet.reset_dimensions()
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, ())
            if not any(_text(value) for value in header):
                continue
            try:
                _column_mapping(header)
            except ValidationError as exc:
                errors.append(exc)
                continue
            if candidates:
                raise ValidationError('В XLSX несколько листов с задачами. Оставьте один лист для импорта.')
            # Missing trailing cells are normal in XLSX; pad to the actual
            # header width, never to the untrusted worksheet dimension.
            padded = (row + (None,) * max(0, len(header) - len(row)) for row in rows)
            candidates.append(_bounded_rows(chain((header,), padded)))
        if not candidates:
            if errors:
                raise errors[0]
            raise ValidationError("В XLSX не найден лист с заголовками задач EVA.")
        if len(candidates) > 1:
            raise ValidationError("В XLSX несколько листов с задачами. Оставьте один лист для импорта.")
        return candidates[0]
    finally:
        workbook.close()


def read_task_file_rows(upload):
    extension = Path(upload.name).suffix.lower()
    if extension not in (".csv", ".xlsx"):
        raise ValidationError("Выберите файл CSV или XLSX.")
    if upload.size > MAX_FILE_BYTES:
        raise ValidationError("Размер файла не должен превышать 10 МБ.")
    data = upload.read(MAX_FILE_BYTES + 1)
    if len(data) > MAX_FILE_BYTES:
        raise ValidationError("Размер файла не должен превышать 10 МБ.")
    try:
        rows = _xlsx_rows(data) if extension == ".xlsx" else _csv_rows(data)
    except (
        BadZipFile, InvalidFileException, KeyError, ValueError,
        OSError, csv.Error, SyntaxError,
    ) as exc:
        raise ValidationError(
            "Не удалось прочитать файл. Повторите выгрузку в CSV или XLSX."
        ) from exc
    if not rows:
        raise ValidationError("Файл пуст.")
    return rows


def task_from_row(row, columns, row_number):
    def value(key):
        index = columns.get(key)
        return _text(row[index]) if index is not None and index < len(row) else ""

    number, title = value("number"), value("title")
    type_name, identifier = value("competency"), value("identifier")
    error = None
    if len(row) <= max(columns.values()):
        error = "недостаточно колонок; проверьте разделитель CSV"
    elif not number or len(number) > 80:
        error = "нужен код задачи длиной до 80 символов"
    elif not title or len(title) > 500:
        error = "нужно название задачи длиной до 500 символов"
    elif _normalise(type_name) not in COMPETENCIES:
        error = f"неизвестный тип «{type_name[:80]}»; укажите аналитику, разработку АБС/BE/FE, тестирование или дефект"
    elif not re.fullmatch(r"CmfTask:[0-9a-fA-F]{8}(?:-[0-9a-fA-F]{4}){3}-[0-9a-fA-F]{12}", identifier):
        error = "нужен идентификатор объекта CmfTask:UUID для ссылки в EVA"
    if len(value("quota")) > 200:
        error = "тип квоты превышает 200 символов"
    if any(len(value(key)) > 500 for key in ("eva_parent_title", "eva_project_name", "eva_epic_title")):
        error = "название родителя, проекта или эпика превышает 500 символов"
    if error:
        raise ValidationError(f"Строка {row_number}: {error}.")
    html = value("html_description")
    description = description_text(html) if html else value("plain_description")
    if len(description) > 100000:
        raise ValidationError(f"Строка {row_number}: описание превышает 100 000 символов.")
    return ImportedTask(
        number=number, title=title, competency=COMPETENCIES[_normalise(type_name)],
        description=description, quota=value("quota"),
        eva_parent_title=value("eva_parent_title"),
        eva_project_name=value("eva_project_name"),
        eva_epic_title=value("eva_epic_title"),
        external_url=f"{EVA_TASK_URL_PREFIX}CmfTask:{UUID(identifier.split(':', 1)[1])}",
    )


def parse_task_file(upload):
    from .sprint_import import parse_sprint_file
    return parse_sprint_file(upload, require_status=False)


def save_task_import(project, parsed):
    from .sprint_import import save_sprint_import
    return save_sprint_import(project, parsed)
