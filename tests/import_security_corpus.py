"""Small synthetic EVA files for bounded parser and HTTP regression checks."""

import csv
from io import BytesIO, StringIO
import struct
import warnings
from xml.sax.saxutils import escape
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile

from openpyxl import Workbook


HEADERS = ['Идентификатор объекта', 'Код', 'Наименование', 'Текст',
           'Логический тип.Имя логического типа', 'Оценка задачи, час', 'Статус.Имя статуса']
ROW = ['CmfTask:5d7adfd4-a87a-11f1-8708-9e913e917b69', 'IMPORT-1',
       'Synthetic import task', '<p>Safe description</p>', 'Разработка АБС', '8', 'В работе']


def csv_file(rows=1):
    stream = StringIO(newline='')
    writer = csv.writer(stream, delimiter=';')
    writer.writerow(HEADERS)
    for index in range(rows):
        writer.writerow([ROW[0][:-8] + f'{index:08x}', 'IMPORT-' + str(index), *ROW[2:]])
    return stream.getvalue().encode('utf-8-sig')


def xlsx_file():
    workbook = Workbook()
    workbook.active.append(HEADERS)
    workbook.active.append(ROW)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def rewrite_xlsx(changes, *, compression=ZIP_DEFLATED):
    stream = BytesIO()
    with ZipFile(BytesIO(xlsx_file())) as source, ZipFile(stream, 'w', compression) as target:
        for item in source.infolist():
            target.writestr(item.filename, changes.get(item.filename, source.read(item)))
        for name, data in changes.items():
            if name not in source.namelist():
                target.writestr(name, data)
    return stream.getvalue()


def unsupported_compression():
    raw = bytearray(xlsx_file())
    with ZipFile(BytesIO(raw)) as source:
        item = source.getinfo('xl/workbook.xml')
        struct.pack_into('<H', raw, item.header_offset + 8, 99)
        offset = source.start_dir
        for entry in source.infolist():
            assert raw[offset:offset + 4] == b'PK\x01\x02'
            if entry.filename == item.filename:
                struct.pack_into('<H', raw, offset + 10, 99)
            filename, extra, comment = struct.unpack_from('<HHH', raw, offset + 28)
            offset += 46 + filename + extra + comment
    return bytes(raw)


def encrypted_member():
    raw = bytearray(xlsx_file())
    with ZipFile(BytesIO(raw)) as source:
        item = source.getinfo('xl/workbook.xml')
        struct.pack_into('<H', raw, item.header_offset + 6, item.flag_bits | 1)
        offset = source.start_dir
        for entry in source.infolist():
            if entry.filename == item.filename:
                struct.pack_into('<H', raw, offset + 8, entry.flag_bits | 1)
            filename, extra, comment = struct.unpack_from('<HHH', raw, offset + 28)
            offset += 46 + filename + extra + comment
    return bytes(raw)


def bad_shared_string():
    return rewrite_xlsx({'xl/worksheets/sheet1.xml':
        b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b'<dimension ref="A1:G2"/><sheetData><row r="1"><c r="A1" t="s">'
        b'<v>999999</v></c></row></sheetData></worksheet>'})


def ambiguous_padded_sheets(sheets=12, rows=5000):
    cells = ''.join(f'<c r="{chr(65 + index)}1" t="inlineStr"><is><t>{escape(value)}</t></is></c>'
                    for index, value in enumerate(HEADERS))
    data = '<row r="1">' + cells + '</row>'
    data += ''.join(f'<row r="{index}"><c r="A{index}" t="inlineStr"><is><t>x</t></is></c></row>'
                    for index in range(2, rows + 1))
    sheet = ('<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
             f'<dimension ref="A1:SR{rows}"/><sheetData>{data}</sheetData></worksheet>').encode()
    workbook = ('<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
                'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets>'
                + ''.join(f'<sheet name="Data{index}" sheetId="{index}" r:id="rId{index}"/>'
                          for index in range(1, sheets + 1)) + '</sheets></workbook>')
    relationships = ('<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        + ''.join(f'<Relationship Id="rId{index}" Type="http://schemas.openxmlformats.org/'
                  f'officeDocument/2006/relationships/worksheet" Target="/xl/worksheets/sheet{index}.xml"/>'
                  for index in range(1, sheets + 1)) + '</Relationships>')
    changes = {'xl/workbook.xml': workbook.encode(), 'xl/_rels/workbook.xml.rels': relationships.encode()}
    changes.update({f'xl/worksheets/sheet{index}.xml': sheet for index in range(1, sheets + 1)})
    return rewrite_xlsx(changes)


def expansion_bomb():
    return rewrite_xlsx({'unused.xml': b' ' * (64 * 1024 * 1024 + 1)})


def duplicate_member():
    stream = BytesIO(xlsx_file())
    with warnings.catch_warnings(), ZipFile(stream, 'a', ZIP_STORED) as archive:
        warnings.simplefilter('ignore', UserWarning)
        archive.writestr('xl/workbook.xml', b'<workbook/>')
    return stream.getvalue()


def xlsx_rows(count):
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet()
    sheet.append(HEADERS)
    for index in range(count):
        sheet.append([ROW[0][:-8] + f'{index:08x}', 'IMPORT-' + str(index), *ROW[2:]])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()
